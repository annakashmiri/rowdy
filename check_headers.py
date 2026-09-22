#!/usr/bin/env python3
"""Check the tool's column layout against the real spreadsheet.

Run this whenever someone adds, moves or renames a column in the sheet:

    python3 tests/check_headers.py path/to/Briefing_Doc.xlsx

Needs: Python 3, openpyxl (pip install openpyxl), Node.
"""
import json, subprocess, sys, pathlib
import openpyxl
from openpyxl.utils import get_column_letter as L

if len(sys.argv) != 2:
    sys.exit(__doc__)

here = pathlib.Path(__file__).parent
specs = json.loads(subprocess.check_output(["node", str(here / "dump-specs.js")]))
wb = openpyxl.load_workbook(sys.argv[1])
tabs = {"meta": "Meta", "tiktok": "TikTok", "youtube": "YouTube", "search": "Search"}
norm = lambda s: "".join(ch for ch in str(s).lower() if ch.isalnum())
problems = 0

for key, cols in specs.items():
    ws = wb[tabs[key]]
    last = max(i for i in range(1, ws.max_column + 1) if ws.cell(4, i).value is not None)
    expected = [L(i) for i in range(1, last + 1)]
    letters = [c["c"] for c in cols]
    if letters != expected:
        print(f"[{tabs[key]}] tool has {len(letters)} columns, sheet header row has {last}. Layout has changed.")
        problems += 1
    for c in cols:
        header = ws[f"{c['c']}4"].value
        if "len" in c:
            if header is None or not ("max" in str(header).lower() or str(header).isdigit()):
                print(f"[{tabs[key]}] {c['c']}4 should be a length column but says {header!r}")
                problems += 1
        else:
            a, b = norm(header), norm(c["n"])
            if not (a in b or b in a):
                print(f"[{tabs[key]}] {c['c']}4: sheet says {header!r}, tool says {c['n']!r}")
                problems += 1

print("All columns match." if not problems else f"\n{problems} problem(s). Update buildSpecs() in index.html.")
sys.exit(1 if problems else 0)
